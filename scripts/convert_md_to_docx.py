"""
Convert Markdown file to Word Document
Usage: python convert_md_to_docx.py
"""

import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def set_cell_shading(cell, color):
    """Set cell background color"""
    shading = OxmlElement('w:shd')
    shading.set(qn('w:fill'), color)
    cell._tc.get_or_add_tcPr().append(shading)

def create_document():
    doc = Document()
    
    # Set up styles
    styles = doc.styles
    
    # Title style
    title_style = styles['Title']
    title_style.font.size = Pt(28)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor(0, 51, 102)
    
    # Heading 1
    h1_style = styles['Heading 1']
    h1_style.font.size = Pt(18)
    h1_style.font.bold = True
    h1_style.font.color.rgb = RGBColor(0, 51, 102)
    
    # Heading 2
    h2_style = styles['Heading 2']
    h2_style.font.size = Pt(14)
    h2_style.font.bold = True
    h2_style.font.color.rgb = RGBColor(0, 102, 153)
    
    # Heading 3
    h3_style = styles['Heading 3']
    h3_style.font.size = Pt(12)
    h3_style.font.bold = True
    h3_style.font.color.rgb = RGBColor(51, 51, 51)
    
    return doc

def add_table(doc, headers, rows):
    """Add a formatted table to the document"""
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Header row
    header_cells = table.rows[0].cells
    for i, header in enumerate(headers):
        header_cells[i].text = header
        header_cells[i].paragraphs[0].runs[0].bold = True
        set_cell_shading(header_cells[i], '0066CC')
        header_cells[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
    
    # Data rows
    for row_data in rows:
        row_cells = table.add_row().cells
        for i, cell_data in enumerate(row_data):
            row_cells[i].text = cell_data
    
    return table

def parse_markdown_table(lines, start_idx):
    """Parse a markdown table and return headers, rows, and end index"""
    headers = []
    rows = []
    i = start_idx
    
    # Parse header
    if i < len(lines) and '|' in lines[i]:
        header_line = lines[i].strip()
        headers = [cell.strip() for cell in header_line.split('|') if cell.strip()]
        i += 1
    
    # Skip separator line
    if i < len(lines) and '|' in lines[i] and '-' in lines[i]:
        i += 1
    
    # Parse rows
    while i < len(lines) and '|' in lines[i]:
        row_line = lines[i].strip()
        if row_line.startswith('|') or '|' in row_line:
            cells = [cell.strip() for cell in row_line.split('|') if cell.strip()]
            if cells:
                rows.append(cells)
        i += 1
    
    return headers, rows, i

def convert_md_to_docx(md_file, docx_file):
    """Convert markdown file to Word document"""
    
    with open(md_file, 'r', encoding='utf-8') as f:
        content = f.read()
    
    doc = create_document()
    lines = content.split('\n')
    
    i = 0
    in_code_block = False
    code_content = []
    
    while i < len(lines):
        line = lines[i]
        
        # Code block handling
        if line.strip().startswith('```'):
            if in_code_block:
                # End code block
                if code_content:
                    code_text = '\n'.join(code_content)
                    p = doc.add_paragraph()
                    run = p.add_run(code_text)
                    run.font.name = 'Consolas'
                    run.font.size = Pt(9)
                    p.paragraph_format.left_indent = Inches(0.3)
                code_content = []
                in_code_block = False
            else:
                # Start code block
                in_code_block = True
            i += 1
            continue
        
        if in_code_block:
            code_content.append(line)
            i += 1
            continue
        
        # Skip empty lines
        if not line.strip():
            i += 1
            continue
        
        # Horizontal rule
        if line.strip() == '---':
            doc.add_paragraph('_' * 50)
            i += 1
            continue
        
        # Headers
        if line.startswith('# '):
            doc.add_heading(line[2:].strip(), 0)
            i += 1
            continue
        elif line.startswith('## '):
            doc.add_heading(line[3:].strip(), 1)
            i += 1
            continue
        elif line.startswith('### '):
            doc.add_heading(line[4:].strip(), 2)
            i += 1
            continue
        elif line.startswith('#### '):
            doc.add_heading(line[5:].strip(), 3)
            i += 1
            continue
        
        # Tables
        if '|' in line and i + 1 < len(lines) and '---' in lines[i + 1]:
            headers, rows, end_idx = parse_markdown_table(lines, i)
            if headers and rows:
                add_table(doc, headers, rows)
                doc.add_paragraph()
            i = end_idx
            continue
        
        # Block quotes
        if line.strip().startswith('>'):
            quote_text = line.strip()[1:].strip().strip('"')
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.5)
            run = p.add_run(quote_text)
            run.italic = True
            run.font.color.rgb = RGBColor(80, 80, 80)
            i += 1
            continue
        
        # Bullet points
        if line.strip().startswith('- ') or line.strip().startswith('* '):
            text = line.strip()[2:]
            # Handle bold text
            text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
            p = doc.add_paragraph(text, style='List Bullet')
            i += 1
            continue
        
        # Numbered lists
        match = re.match(r'^(\d+)\.\s+(.+)$', line.strip())
        if match:
            text = match.group(2)
            text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
            p = doc.add_paragraph(text, style='List Number')
            i += 1
            continue
        
        # Regular paragraph
        text = line.strip()
        if text:
            # Remove markdown formatting
            text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
            text = re.sub(r'\*(.*?)\*', r'\1', text)
            text = re.sub(r'`(.*?)`', r'\1', text)
            text = re.sub(r'\[(.*?)\]\(.*?\)', r'\1', text)
            
            # Skip pure emoji lines
            if text and not all(c in '✅🔄📊🎬🎯📁📞ℹ️🔧📋' for c in text.replace(' ', '')):
                p = doc.add_paragraph(text)
        
        i += 1
    
    doc.save(docx_file)
    print(f"Document saved to: {docx_file}")

if __name__ == "__main__":
    md_file = r"C:\trainings\factoryAI\SDLC_Demo\docs\Factory_AI_Lineaje_Secure_SDLC_Demo.md"
    docx_file = r"C:\trainings\factoryAI\SDLC_Demo\docs\Factory_AI_Lineaje_Secure_SDLC_Demo.docx"
    
    convert_md_to_docx(md_file, docx_file)
