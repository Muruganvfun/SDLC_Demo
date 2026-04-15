"""
Convert Interactive Demo Script Markdown to Word Document
Usage: python create_interactive_demo_docx.py
"""

import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Twips
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def set_cell_shading(cell, color):
    """Set cell background color"""
    shading = OxmlElement('w:shd')
    shading.set(qn('w:fill'), color)
    cell._tc.get_or_add_tcPr().append(shading)

def add_speech_box(doc, text, speaker_note=False):
    """Add a speech/presenter notes box"""
    # Create a table with one cell to simulate a box
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    cell = table.rows[0].cells[0]
    
    # Set background color
    if speaker_note:
        set_cell_shading(cell, 'FFF3CD')  # Light yellow for speaker notes
    else:
        set_cell_shading(cell, 'E8F4FD')  # Light blue for speech
    
    # Add the text
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(8)
    
    # Add icon
    if speaker_note:
        run = paragraph.add_run("[Presenter Note] ")
        run.bold = True
        run.font.color.rgb = RGBColor(133, 100, 4)
    else:
        run = paragraph.add_run("[SPEECH] ")
        run.bold = True
        run.font.color.rgb = RGBColor(13, 110, 253)
    
    # Add speech text
    run = paragraph.add_run(text)
    run.italic = True
    run.font.size = Pt(11)
    
    doc.add_paragraph()

def add_action_box(doc, text):
    """Add a demo action box"""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    cell = table.rows[0].cells[0]
    set_cell_shading(cell, 'D4EDDA')  # Light green for actions
    
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(8)
    
    run = paragraph.add_run("[ACTION] ")
    run.bold = True
    run.font.color.rgb = RGBColor(25, 135, 84)
    
    run = paragraph.add_run(text)
    run.font.size = Pt(11)
    
    doc.add_paragraph()

def add_code_block(doc, code, language=""):
    """Add a formatted code block"""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    cell = table.rows[0].cells[0]
    set_cell_shading(cell, '1E1E1E')  # Dark background like VS Code
    
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.left_indent = Inches(0.1)
    
    run = paragraph.add_run(code)
    run.font.name = 'Consolas'
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(212, 212, 212)  # Light gray text
    
    doc.add_paragraph()

def add_timing_badge(doc, timing):
    """Add a timing badge"""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run(f"[{timing}]")
    run.bold = True
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(108, 117, 125)

def add_formatted_table(doc, headers, rows):
    """Add a formatted table"""
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Header row
    header_cells = table.rows[0].cells
    for i, header in enumerate(headers):
        header_cells[i].text = header
        header_cells[i].paragraphs[0].runs[0].bold = True
        set_cell_shading(header_cells[i], '0066CC')
        header_cells[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
    
    # Data rows
    for row_data in rows:
        row_cells = table.add_row().cells
        for i, cell_data in enumerate(row_data):
            if i < len(row_cells):
                row_cells[i].text = str(cell_data)
    
    doc.add_paragraph()

def create_document():
    """Create document with custom styles"""
    doc = Document()
    
    # Set up page margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
    
    # Customize styles
    styles = doc.styles
    
    # Title
    title_style = styles['Title']
    title_style.font.size = Pt(28)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor(0, 51, 102)
    
    # Heading 1
    h1_style = styles['Heading 1']
    h1_style.font.size = Pt(20)
    h1_style.font.bold = True
    h1_style.font.color.rgb = RGBColor(0, 51, 102)
    
    # Heading 2
    h2_style = styles['Heading 2']
    h2_style.font.size = Pt(16)
    h2_style.font.bold = True
    h2_style.font.color.rgb = RGBColor(0, 102, 153)
    
    # Heading 3
    h3_style = styles['Heading 3']
    h3_style.font.size = Pt(14)
    h3_style.font.bold = True
    h3_style.font.color.rgb = RGBColor(51, 51, 51)
    
    return doc

def process_markdown(md_content, doc):
    """Process markdown content and add to document"""
    lines = md_content.split('\n')
    
    i = 0
    in_code_block = False
    code_content = []
    code_language = ""
    in_quote = False
    quote_content = []
    
    while i < len(lines):
        line = lines[i]
        
        # Code block handling
        if line.strip().startswith('```'):
            if in_code_block:
                # End code block
                if code_content:
                    add_code_block(doc, '\n'.join(code_content), code_language)
                code_content = []
                code_language = ""
                in_code_block = False
            else:
                # Start code block
                code_language = line.strip()[3:].strip()
                in_code_block = True
            i += 1
            continue
        
        if in_code_block:
            code_content.append(line)
            i += 1
            continue
        
        # Handle multi-line quotes
        if line.strip().startswith('>'):
            quote_text = line.strip()[1:].strip()
            
            # Check if this is a speech indicator
            if quote_text.startswith('**[') and ']**' in quote_text:
                # Speaker direction
                if quote_content:
                    add_speech_box(doc, ' '.join(quote_content))
                    quote_content = []
                
                direction = quote_text.replace('**', '').strip()
                p = doc.add_paragraph()
                run = p.add_run(direction)
                run.bold = True
                run.italic = True
                run.font.color.rgb = RGBColor(108, 117, 125)
                i += 1
                continue
            
            # Regular speech content
            clean_text = quote_text.strip('"').strip()
            if clean_text:
                quote_content.append(clean_text)
            i += 1
            continue
        elif quote_content:
            # End of quote block
            add_speech_box(doc, ' '.join(quote_content))
            quote_content = []
        
        # Skip empty lines
        if not line.strip():
            i += 1
            continue
        
        # Horizontal rule
        if line.strip() == '---':
            doc.add_paragraph()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run('━' * 40)
            run.font.color.rgb = RGBColor(200, 200, 200)
            doc.add_paragraph()
            i += 1
            continue
        
        # Headers
        if line.startswith('# '):
            doc.add_heading(line[2:].strip(), 0)
            i += 1
            continue
        elif line.startswith('## '):
            doc.add_heading(line[3:].strip(), 1)
            i += 1
            continue
        elif line.startswith('### '):
            # Check if this is a "Visual" or "Speech Content" header
            header_text = line[4:].strip()
            if header_text == 'Visual':
                p = doc.add_paragraph()
                run = p.add_run('[Visual Setup]')
                run.bold = True
                run.font.color.rgb = RGBColor(156, 39, 176)
                i += 1
                continue
            elif header_text == 'Speech Content':
                p = doc.add_paragraph()
                run = p.add_run('[What to Say]')
                run.bold = True
                run.font.color.rgb = RGBColor(0, 150, 136)
                i += 1
                continue
            elif 'Demo Action' in header_text or 'Live Demo' in header_text:
                p = doc.add_paragraph()
                run = p.add_run(f'[{header_text}]')
                run.bold = True
                run.font.color.rgb = RGBColor(25, 135, 84)
                i += 1
                continue
            elif header_text == 'Transition':
                p = doc.add_paragraph()
                run = p.add_run('[Transition]')
                run.bold = True
                run.font.color.rgb = RGBColor(255, 152, 0)
                i += 1
                continue
            else:
                doc.add_heading(header_text, 2)
            i += 1
            continue
        elif line.startswith('#### '):
            doc.add_heading(line[5:].strip(), 3)
            i += 1
            continue
        
        # Tables
        if '|' in line and i + 1 < len(lines) and '---' in lines[i + 1]:
            headers = [cell.strip() for cell in line.split('|') if cell.strip()]
            i += 2  # Skip header and separator
            rows = []
            while i < len(lines) and '|' in lines[i]:
                cells = [cell.strip() for cell in lines[i].split('|') if cell.strip()]
                if cells:
                    rows.append(cells)
                i += 1
            if headers and rows:
                add_formatted_table(doc, headers, rows)
            continue
        
        # Special markers for demo
        if line.strip().startswith('**Demo Action:**') or line.strip().startswith('Demo Action:'):
            text = line.strip().replace('**Demo Action:**', '').replace('Demo Action:', '').strip()
            add_action_box(doc, text)
            i += 1
            continue
        
        # Check for timing markers like "(30 seconds)" or "(3 minutes)"
        timing_match = re.search(r'\((\d+\s*(?:seconds?|minutes?))\)', line)
        if timing_match:
            add_timing_badge(doc, timing_match.group(1))
            line = line.replace(timing_match.group(0), '').strip()
        
        # Bullet points
        if line.strip().startswith('- ') or line.strip().startswith('* '):
            text = line.strip()[2:]
            # Handle bold text within bullets
            text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
            p = doc.add_paragraph(text, style='List Bullet')
            i += 1
            continue
        
        # Numbered lists
        match = re.match(r'^(\d+)\.\s+(.+)$', line.strip())
        if match:
            text = match.group(2)
            text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
            p = doc.add_paragraph(text, style='List Number')
            i += 1
            continue
        
        # Regular paragraph
        text = line.strip()
        if text:
            # Clean up formatting
            text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
            text = re.sub(r'\*(.*?)\*', r'\1', text)
            text = re.sub(r'`(.*?)`', r'\1', text)
            text = re.sub(r'\[(.*?)\]\(.*?\)', r'\1', text)
            
            # Add paragraph
            p = doc.add_paragraph(text)
        
        i += 1
    
    # Handle any remaining quote content
    if quote_content:
        add_speech_box(doc, ' '.join(quote_content))

def main():
    md_file = r"C:\trainings\factoryAI\SDLC_Demo\docs\Factory_AI_Lineaje_Interactive_Demo_Script.md"
    docx_file = r"C:\trainings\factoryAI\SDLC_Demo\docs\Factory_AI_Lineaje_Interactive_Demo_Script.docx"
    
    # Read markdown
    with open(md_file, 'r', encoding='utf-8') as f:
        md_content = f.read()
    
    # Create document
    doc = create_document()
    
    # Process and add content
    process_markdown(md_content, doc)
    
    # Save
    doc.save(docx_file)
    print(f"Interactive demo script saved to: {docx_file}")

if __name__ == "__main__":
    main()
